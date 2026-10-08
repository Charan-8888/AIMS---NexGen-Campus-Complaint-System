import sys
from pptx import Presentation

def read_presentation(file_path):
    prs = Presentation(file_path)
    with open('ppt_contents.txt', 'w', encoding='utf-8') as f:
        for i, slide in enumerate(prs.slides):
            f.write(f"--- Slide {i} ---\n")
            for j, shape in enumerate(slide.shapes):
                if shape.has_text_frame:
                    text = shape.text.strip()
                    f.write(f"Shape {j} ({shape.name}) TEXT: {repr(text)}\n")
                elif shape.has_table:
                    f.write(f"Shape {j} ({shape.name}) TABLE:\n")
                    for row in shape.table.rows:
                        row_data = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
                        f.write(f"  | {' | '.join(row_data)} |\n")

if __name__ == '__main__':
    read_presentation('AIMS_Campus_Review_4.pptx')
